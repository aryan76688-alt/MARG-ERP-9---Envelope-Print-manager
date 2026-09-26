import io
import re
from typing import Dict, List, Any, Optional, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Column alias dictionary for MARG ERP and Indian accounting software
MARG_COLUMN_ALIASES: Dict[str, List[str]] = {
    "party_name": [
        "party name", "party", "account name", "ledger name", "partyname",
        "customer name", "customer", "party_name", "account", "client name",
        "firm name", "party title", "name", "cust name", "parties",
        "particulars", "debtor", "sundry debtors", "client", "firm", "ledger",
        "a/c name", "a/c", "ac name", "ac_name", "party/ledger", "dr name", "dr. name"
    ],
    "party_code": [
        "party code", "code", "account code", "ledger code", "partycode",
        "cust code", "customer code", "ac code", "ac_code", "pcode", "id",
        "ledger id", "party id", "sr no", "sr.no", "s.no", "sno", "serial no", "ac no", "acc no"
    ],
    "route": [
        "route", "route name", "delivery route", "beat", "area route", "route no", "routeno",
        "delivery area", "transport route", "zone", "group", "route / area", "station / route"
    ],
    "address": [
        "address", "address 1", "address1", "address_1", "address-1", "party address", "street",
        "address line 1", "address line1", "addressline1", "addr1", "addr 1", "addr_1",
        "full address", "location", "premises", "street 1", "add1", "add 1", "add_1",
        "address (line 1)", "line 1", "line1", "street1", "addr. 1", "addr."
    ],
    "address_line_2": [
        "address line 2", "address line2", "addressline2", "address 2", "address2", "address_2", "address-2",
        "address.1", "addr2", "addr 2", "addr_2", "line 2", "line2", "street 2", "street2",
        "area", "colony", "add2", "add 2", "add_2", "address (line 2)", "addr. 2",
        "locality", "mohalla", "society", "road"
    ],
    "address_line_3": [
        "address line 3", "address line3", "addressline3", "address 3", "address3", "address_3", "address-3",
        "address.2", "addr3", "addr 3", "addr_3", "line 3", "line3", "street 3", "street3",
        "landmark", "near", "add3", "add 3", "add_3", "address (line 3)", "addr. 3",
        "opp", "opposite", "behind", "building", "bldg"
    ],
    "city": [
        "city", "town", "station", "place", "city / town", "district", "dist", "taluka", "headquarter"
    ],
    "state": [
        "state", "state name", "province", "region", "st"
    ],
    "mobile_no": [
        "mobile no", "mobile", "mobile number", "mobile no.", "contact no",
        "contact number", "cell", "phone", "phone no", "phone no.", "cell no",
        "ph no", "ph", "contact", "whatsapp", "mobile 1", "phone 1"
    ],
    "landline": [
        "landline", "landline no", "tel", "telephone", "office phone", "phone (o)",
        "landline number", "phone 2", "office tel"
    ],
    "email": [
        "email", "e-mail", "email id", "email address", "mail"
    ],
    "gst_no": [
        "gst no", "gst", "gst number", "gstin", "gstin/uin", "tax id", "tin"
    ],
    "notes": [
        "notes", "remark", "remarks", "dispatch instruction", "comment", "description", "narration"
    ]
}

def clean_header(h: Any) -> str:
    if h is None:
        return ""
    return str(h).strip().lower()

def detect_column_mappings(headers: List[str]) -> Dict[str, Optional[str]]:
    """
    Automatically detects best-matching column headers for MARG ERP party imports.
    Supports Route and all 3 address lines (address, address_line_2, address_line_3).
    Strictly ignores and excludes any PIN code / Postal code columns.
    """
    mapping: Dict[str, Optional[str]] = {
        "party_name": None,
        "party_code": None,
        "route": None,
        "address": None,
        "address_line_2": None,
        "address_line_3": None,
        "city": None,
        "state": None,
        "mobile_no": None,
        "landline": None,
        "email": None,
        "gst_no": None,
        "notes": None
    }
    
    used_headers = set()
    cleaned = [(orig, clean_header(orig)) for orig in headers if orig]

    def is_pin(text: str) -> bool:
        return any(k in text for k in ["pin", "zip", "postal", "pincode"])

    # 1. Exact alias matching (normalized)
    for field, aliases in MARG_COLUMN_ALIASES.items():
        if mapping[field] is not None:
            continue
        for orig, cl in cleaned:
            if orig in used_headers or is_pin(cl):
                continue
            norm_cl = re.sub(r"[^a-z0-9]", "", cl)
            for alias in aliases:
                norm_alias = re.sub(r"[^a-z0-9]", "", alias)
                if cl == alias or norm_cl == norm_alias:
                    mapping[field] = orig
                    used_headers.add(orig)
                    break
            if mapping[field] is not None:
                break

    # 2. Word-boundary / token matching (aliases >= 3 characters)
    for field, aliases in MARG_COLUMN_ALIASES.items():
        if mapping[field] is not None:
            continue
        for orig, cl in cleaned:
            if orig in used_headers or is_pin(cl):
                continue
            for alias in aliases:
                if len(alias) >= 3:
                    pattern = r'\b' + re.escape(alias) + r'\b'
                    if re.search(pattern, cl):
                        mapping[field] = orig
                        used_headers.add(orig)
                        break
            if mapping[field] is not None:
                break

    # 3. Substring matching for longer keywords (>= 4 chars)
    for field, aliases in MARG_COLUMN_ALIASES.items():
        if mapping[field] is not None:
            continue
        for orig, cl in cleaned:
            if orig in used_headers or is_pin(cl):
                continue
            for alias in aliases:
                if len(alias) >= 4 and (alias in cl or cl in alias):
                    mapping[field] = orig
                    used_headers.add(orig)
                    break
            if mapping[field] is not None:
                break

    # 4. Sequential 3-Address Line Fallback:
    unmapped_addr_cols = [
        orig for orig, cl in cleaned 
        if orig not in used_headers and ("address" in cl or "addr" in cl or "street" in cl or "line" in cl) and not is_pin(cl)
    ]
    if mapping["address"] is None and unmapped_addr_cols:
        mapping["address"] = unmapped_addr_cols.pop(0)
        used_headers.add(mapping["address"])
    if mapping["address_line_2"] is None and unmapped_addr_cols:
        mapping["address_line_2"] = unmapped_addr_cols.pop(0)
        used_headers.add(mapping["address_line_2"])
    if mapping["address_line_3"] is None and unmapped_addr_cols:
        mapping["address_line_3"] = unmapped_addr_cols.pop(0)
        used_headers.add(mapping["address_line_3"])

    # 5. Last resort fallback for party_name:
    if mapping["party_name"] is None:
        for orig, cl in cleaned:
            if orig not in used_headers and not is_pin(cl):
                if any(k in cl for k in ["name", "party", "ledger", "account", "customer"]):
                    mapping["party_name"] = orig
                    used_headers.add(orig)
                    break

    # 6. Fallback for route:
    if mapping["route"] is None:
        for orig, cl in cleaned:
            if orig not in used_headers and not is_pin(cl):
                if "route" in cl or "beat" in cl:
                    mapping["route"] = orig
                    used_headers.add(orig)
                    break

    return mapping

def _load_raw_table(file_bytes: bytes) -> Tuple[List[List[Any]], str]:
    """
    Safely loads tabular data from diverse formats produced by MARG ERP 9+:
    1. openpyxl (Modern .xlsx)
    2. xlrd (Legacy binary .xls BIFF8)
    3. pandas.read_html (HTML tables exported with .xls extension)
    4. pandas.read_csv (CSV / TSV text exports)
    """
    # 1. Try openpyxl (.xlsx)
    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(file_bytes),
            data_only=True,
            read_only=True,
            keep_vba=False
        )
        sheet = wb.active
        sheet_name = sheet.title or "Sheet1"
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        wb.close()
        if rows and any(any(c is not None for c in r) for r in rows):
            return rows, sheet_name
    except Exception:
        pass

    # 2. Try xlrd (binary .xls)
    try:
        import xlrd
        wb = xlrd.open_workbook(file_contents=file_bytes)
        sheet = wb.sheet_by_index(0)
        sheet_name = sheet.name or "Sheet1"
        rows = []
        for r in range(sheet.nrows):
            rows.append([sheet.cell_value(r, c) for c in range(sheet.ncols)])
        if rows and any(any(c not in (None, "") for c in r) for r in rows):
            return rows, sheet_name
    except Exception:
        pass

    # 3. Try pandas read_html (HTML disguised as .xls)
    try:
        import pandas as pd
        dfs = pd.read_html(io.BytesIO(file_bytes))
        if dfs:
            df = dfs[0]
            headers = [str(c) for c in df.columns]
            rows = [headers] + [[cell if pd.notna(cell) else "" for cell in r] for r in df.values.tolist()]
            return rows, "HTML_Export"
    except Exception:
        pass

    # 4. Try CSV / TSV
    try:
        import pandas as pd
        for sep in [",", "\t", ";"]:
            try:
                df = pd.read_csv(io.BytesIO(file_bytes), sep=sep, encoding="utf-8-sig", on_bad_lines="skip")
                if not df.empty and len(df.columns) >= 2:
                    headers = [str(c) for c in df.columns]
                    rows = [headers] + [[cell if pd.notna(cell) else "" for cell in r] for r in df.values.tolist()]
                    return rows, "CSV_Export"
            except Exception:
                continue
    except Exception:
        pass

    raise ValueError("Unable to read Excel workbook. Supported formats: .xlsx, .xls, and MARG ERP HTML/CSV exports.")

def read_excel_file(file_bytes: bytes) -> Tuple[List[str], List[Dict[str, Any]], str, int]:
    """
    Safely reads an Excel workbook without executing formulas or macros.
    Automatically detects the true header row even if title rows precede it.
    Returns (headers, rows, sheet_name, total_rows).
    """
    raw_table, sheet_name = _load_raw_table(file_bytes)

    if not raw_table:
        return [], [], sheet_name, 0

    # Scan the first 15 rows to find the true column header row
    HEADER_KEYWORDS = {
        "party", "name", "ledger", "account", "customer", "code", "address",
        "addr", "add1", "add2", "add3", "city", "station", "state", "mobile",
        "phone", "contact", "route", "area", "gst", "gstin", "email", "remarks",
        "notes", "sno", "s.no", "sr", "particulars", "ac", "dr", "balance"
    }

    best_header_idx = 0
    max_score = 0

    scan_limit = min(len(raw_table), 15)
    for r_idx in range(scan_limit):
        row = raw_table[r_idx]
        if not row or not any(row):
            continue
        score = 0
        for cell in row:
            if cell is None:
                continue
            cell_str = str(cell).strip().lower()
            cell_norm = re.sub(r"[^a-z0-9]", "", cell_str)
            if any(k in cell_norm for k in HEADER_KEYWORDS):
                score += 1
        if score > max_score:
            max_score = score
            best_header_idx = r_idx

    # If no row had keyword matches, pick first row with at least 2 non-empty cells
    if max_score < 2:
        for r_idx in range(scan_limit):
            row = raw_table[r_idx]
            non_empty = [c for c in row if c is not None and str(c).strip()]
            if len(non_empty) >= 2:
                best_header_idx = r_idx
                break

    # Build unique headers
    raw_headers = raw_table[best_header_idx]
    headers: List[str] = []
    seen_headers = set()

    for col_idx, cell in enumerate(raw_headers):
        if cell is not None and str(cell).strip():
            h_text = str(cell).strip()
        else:
            h_text = f"Column_{col_idx + 1}"
        
        # Deduplicate identical header names (e.g. "Address", "Address")
        base_h = h_text
        dup_count = 2
        while h_text in seen_headers:
            h_text = f"{base_h}_{dup_count}"
            dup_count += 1
        seen_headers.add(h_text)
        headers.append(h_text)

    # Process data rows
    rows: List[Dict[str, Any]] = []
    for row in raw_table[best_header_idx + 1:]:
        if not row or not any(row):
            continue

        if all(c is None or str(c).strip() == "" for c in row):
            continue

        row_dict: Dict[str, Any] = {}
        has_any_val = False
        for col_idx, cell_value in enumerate(row):
            if col_idx < len(headers):
                h = headers[col_idx]
                if cell_value is not None:
                    val_str = str(cell_value).strip()
                    # Remove trailing .0 from float numbers
                    if isinstance(cell_value, float) and cell_value.is_integer():
                        val_str = str(int(cell_value))
                    elif isinstance(cell_value, float) and val_str.endswith(".0"):
                        val_str = val_str[:-2]
                    row_dict[h] = val_str
                    if val_str:
                        has_any_val = True
                else:
                    row_dict[h] = ""
        
        if has_any_val:
            rows.append(row_dict)

    return headers, rows, sheet_name, len(rows)

def validate_imported_rows(
    raw_rows: List[Dict[str, Any]],
    mapping: Dict[str, Optional[str]],
    existing_party_names: set,
    existing_party_codes: set
) -> Dict[str, Any]:
    """
    Validates rows against MARG rules and existing DB records.
    Returns categorized rows: valid_rows, warning_rows, error_rows.
    Guarantees that rows with a party name are never rejected due to missing address lines;
    they are intelligently auto-filled with warnings instead of fatal errors.
    """
    valid_rows = []
    warning_rows = []
    error_rows = []

    seen_names_in_file = set()
    seen_codes_in_file = set()

    for idx, row in enumerate(raw_rows, start=1):
        def get_val(key: str) -> str:
            col = mapping.get(key)
            if col and col in row:
                return str(row[col]).strip()
            return ""

        party_name = get_val("party_name")
        party_code = get_val("party_code")
        route = get_val("route")
        address = get_val("address")
        address_line_2 = get_val("address_line_2")
        address_line_3 = get_val("address_line_3")
        city = get_val("city")
        state = get_val("state")
        mobile = get_val("mobile_no")
        landline = get_val("landline")
        email = get_val("email")
        gst_no = get_val("gst_no")
        notes = get_val("notes")

        # Normalize mobile (strip non-digits)
        clean_mob = re.sub(r"[^\d]", "", mobile) if mobile else ""

        errors = []
        warnings = []

        # Party name is the only strictly fatal requirement
        if not party_name:
            errors.append("Party Name is required")

        # Intelligent address cascade:
        # If address (line 1) is empty, pull from line 2 or line 3 or city
        if not address:
            if address_line_2:
                address = address_line_2
                address_line_2 = address_line_3
                address_line_3 = ""
            elif address_line_3:
                address = address_line_3
                address_line_3 = ""
            elif city:
                address = city
                warnings.append("Address Line 1 was empty; defaulted to City")
            else:
                address = "DAHEGAM"
                warnings.append("Address Line 1 missing; defaulted to DAHEGAM")

        # City / State defaults if missing
        if not city:
            warnings.append("City is missing (defaults to DAHEGAM)")
            city = "DAHEGAM"
        if not state:
            warnings.append("State is missing (defaults to GUJARAT)")
            state = "GUJARAT"

        # Duplicate checking against Database & File
        norm_name = party_name.strip().upper()
        norm_code = party_code.strip().upper() if party_code else ""
        is_already_exists = False
        match_reason = ""

        if norm_name and norm_name in existing_party_names:
            is_already_exists = True
            match_reason = f"Party name '{party_name}' already exists in ledger"
            warnings.append(match_reason)
        elif norm_code and norm_code in existing_party_codes:
            is_already_exists = True
            match_reason = f"Party code '{party_code}' already exists in ledger"
            warnings.append(match_reason)

        if norm_name:
            if norm_name in seen_names_in_file:
                warnings.append(f"Party name '{party_name}' repeated multiple times in this Excel file")
            seen_names_in_file.add(norm_name)

        if norm_code:
            if norm_code in seen_codes_in_file:
                warnings.append(f"Party code '{party_code}' repeated in this Excel file")
            seen_codes_in_file.add(norm_code)

        # Mobile validation (Indian 10-digit mobile)
        if mobile:
            if len(clean_mob) < 10:
                warnings.append(f"Mobile number '{mobile}' seems incomplete")

        # Email validation
        if email and "@" not in email:
            warnings.append(f"Email '{email}' format invalid")

        row_status = "error" if errors else ("already_exists" if is_already_exists else ("warning" if warnings else "valid"))

        record = {
            "row_index": idx,
            "party_name": party_name,
            "party_code": party_code,
            "route": route,
            "address": address,
            "address_line_2": address_line_2,
            "address_line_3": address_line_3,
            "city": city,
            "state": state,
            "mobile_no": mobile,
            "landline": landline,
            "email": email,
            "gst_no": gst_no,
            "notes": notes,
            "errors": errors,
            "warnings": warnings,
            "is_already_exists": is_already_exists,
            "match_reason": match_reason,
            "status": row_status
        }

        if errors:
            error_rows.append(record)
        elif is_already_exists:
            warning_rows.append(record)
        elif warnings:
            warning_rows.append(record)
        else:
            valid_rows.append(record)

    already_exists_rows = [r for r in warning_rows if r.get("is_already_exists")]
    new_rows = valid_rows + [r for r in warning_rows if not r.get("is_already_exists")]

    return {
        "total_rows": len(raw_rows),
        "valid_count": len(valid_rows),
        "warning_count": len(warning_rows),
        "error_count": len(error_rows),
        "already_exists_count": len(already_exists_rows),
        "new_count": len(new_rows),
        "valid_rows": valid_rows,
        "warning_rows": warning_rows,
        "error_rows": error_rows,
        "already_exists_rows": already_exists_rows,
        "all_preview_rows": valid_rows + warning_rows + error_rows
    }

def generate_sample_excel_template() -> bytes:
    """
    Creates a pre-formatted Excel workbook template for office staff to import MARG ERP party data.
    Includes Delivery Route and 3 Address lines. Strictly excludes PIN code.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MARG Parties"

    # Header styling: MARG navy header with bold white text
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    columns = [
        ("Party Name *", 28),
        ("Party Code", 14),
        ("Delivery Route", 18),
        ("Address Line 1 *", 35),
        ("Address Line 2", 30),
        ("Address Line 3", 25),
        ("City *", 18),
        ("State *", 18),
        ("Mobile No.", 16),
        ("Landline", 16),
        ("Email", 25),
        ("GST No.", 20),
        ("Notes", 30)
    ]

    for col_idx, (col_name, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[1].height = 28

    # Realistic MARG ERP sample data rows with Delivery Route and 3 Address Lines
    sample_data = [
        ["JODHPUR MEDICOSE", "P0001", "RAJASTHAN ROUTE", "SHREE MOHANGADH", "NEAR STN ROAD", "OPP CIVIL HOSPITAL", "JODHPUR", "RAJASTHAN", "9829012345", "0291-2645120", "jodhpurmedicose@gmail.com", "08ABCDE1234F1Z2", "Express courier"],
        ["JAY SHREE TRADERS", "P0002", "CITY MAIN ROUTE", "SHOP 14, APMC MARKET", "SECTOR 19", "PHARMA WING", "AHMEDABAD", "GUJARAT", "9876543210", "079-25418900", "jayshreetraders@yahoo.com", "24ABCDE5678G2Z1", "Medical consignments"],
        ["JIGNESH ENTERPRISE", "P0003", "VADODARA HIGHWAY", "GIDC PHASE 2", "PLOT 45/A", "", "VADODARA", "GUJARAT", "9824098765", "0265-2890123", "jignesh_ent@rediffmail.com", "24XYZAB9876C1Z8", "Priority"],
        ["J K PHARMA DISTRIBUTOR", "P0004", "SOUTH GUJARAT", "MEDICINE COMPLEX", "RING ROAD", "", "SURAT", "GUJARAT", "9898011223", "0261-2478901", "jkpharma@suratpharma.com", "24LMNOP4321D1Z9", ""],
        ["JALARAM AGENCIES", "P0005", "SAURASHTRA ROUTE", "GRAIN MARKET", "OPP TOWN HALL", "", "RAJKOT", "GUJARAT", "9426033445", "0281-2234567", "jalaram_rajkot@gmail.com", "24PQRSU8765E1Z4", ""],
        ["JANTA MEDICALS", "P0006", "NORTH ZONE", "MAIN HOSPITAL ROAD", "NEAR BUS STAND", "", "JAIPUR", "RAJASTHAN", "9829567890", "0141-2356789", "jantamedicals.jpr@gmail.com", "08JKLMN3456H1Z3", "Cold chain parcel"]
    ]

    for row_idx, row_values in enumerate(sample_data, start=2):
        ws.row_dimensions[row_idx].height = 20
        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [2, 3, 9, 10]:
                cell.alignment = Alignment(horizontal="center")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

def export_parties_to_excel(parties: List[Any]) -> bytes:
    """Exports parties to a clean, formatted XLSX spreadsheet with separate English and Gujarati sheets."""
    wb = openpyxl.Workbook()

    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=10)

    # ==========================================
    # SHEET 1: ENGLISH
    # ==========================================
    ws_en = wb.active
    ws_en.title = "English"
    en_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

    en_headers = [
        ("Party Name", 28),
        ("Code", 14),
        ("Delivery Route", 18),
        ("Address Line 1", 35),
        ("Address Line 2", 30),
        ("Address Line 3", 25),
        ("City", 18),
        ("State", 18),
        ("Mobile No.", 16),
        ("Landline", 16),
        ("Email", 25),
        ("GST No.", 20),
        ("Status", 12)
    ]

    for col_idx, (col_name, width) in enumerate(en_headers, start=1):
        cell = ws_en.cell(row=1, column=col_idx, value=col_name)
        cell.fill = en_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        col_letter = get_column_letter(col_idx)
        ws_en.column_dimensions[col_letter].width = width

    ws_en.row_dimensions[1].height = 26

    for row_idx, p in enumerate(parties, start=2):
        ws_en.row_dimensions[row_idx].height = 20
        row_vals = [
            p.party_name,
            p.party_code or "",
            getattr(p, "route", "") or "",
            p.address,
            getattr(p, "address_line_2", "") or "",
            getattr(p, "address_line_3", "") or "",
            p.city,
            p.state,
            p.mobile_no or "",
            p.landline or "",
            p.email or "",
            p.gst_no or "",
            "Active" if p.is_active else "Inactive"
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_en.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border

    # ==========================================
    # SHEET 2: GUJARATI (ગુજરાતી)
    # ==========================================
    ws_gu = wb.create_sheet(title="Gujarati")
    gu_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid") # Deep Emerald

    gu_headers = [
        ("પાર્ટીનું નામ (Party Name)", 30),
        ("કોડ (Code)", 14),
        ("ડિલિવરી રૂટ (Route)", 18),
        ("સરનામું ૧ (Address Line 1)", 35),
        ("સરનામું ૨ (Address Line 2)", 30),
        ("સરનામું ૩ (Address Line 3)", 25),
        ("શહેર (City)", 18),
        ("રાજ્ય (State)", 18),
        ("મોબાઈલ નંબર (Mobile)", 16),
        ("લેન્ડલાઈન (Phone)", 16),
        ("ઇમેઇલ (Email)", 25),
        ("GST નંબર (GST No.)", 20),
        ("સ્થિતિ (Status)", 12)
    ]

    for col_idx, (col_name, width) in enumerate(gu_headers, start=1):
        cell = ws_gu.cell(row=1, column=col_idx, value=col_name)
        cell.fill = gu_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        col_letter = get_column_letter(col_idx)
        ws_gu.column_dimensions[col_letter].width = width

    ws_gu.row_dimensions[1].height = 26

    for row_idx, p in enumerate(parties, start=2):
        ws_gu.row_dimensions[row_idx].height = 20
        # Use Gujarati translation fields if present, else fallback gracefully
        gu_name = getattr(p, "party_name_gu", None) or p.party_name
        gu_addr = getattr(p, "address_gu", None) or p.address
        gu_addr2 = getattr(p, "address_line_2_gu", None) or getattr(p, "address_line_2", "") or ""
        gu_addr3 = getattr(p, "address_line_3_gu", None) or getattr(p, "address_line_3", "") or ""
        gu_city = getattr(p, "city_gu", None) or p.city
        gu_state = getattr(p, "state_gu", None) or p.state

        row_vals_gu = [
            gu_name,
            p.party_code or "",
            getattr(p, "route", "") or "",
            gu_addr,
            gu_addr2,
            gu_addr3,
            gu_city,
            gu_state,
            p.mobile_no or "",
            p.landline or "",
            p.email or "",
            p.gst_no or "",
            "ચાલુ (Active)" if p.is_active else "બંધ (Inactive)"
        ]
        for col_idx, val in enumerate(row_vals_gu, start=1):
            cell = ws_gu.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

def export_print_history_to_excel(jobs: List[Any]) -> bytes:
    """Exports print jobs history to an XLSX spreadsheet with user attribution."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Print History"

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    headers = [
        ("Job No.", 18),
        ("Date & Time", 20),
        ("Created By", 16),
        ("Party Name", 28),
        ("Parcel Type", 16),
        ("Cases", 10),
        ("Total Weight", 14),
        ("Size", 12),
        ("Printer", 24),
        ("Status", 12)
    ]

    for col_idx, (col_name, width) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[1].height = 26

    for row_idx, j in enumerate(jobs, start=2):
        ws.row_dimensions[row_idx].height = 20
        date_str = j.created_at.strftime("%d-%m-%Y %H:%M") if j.created_at else ""
        creator = getattr(j, "created_by", "Admin") or "Admin"
        row_vals = [
            j.job_number,
            date_str,
            creator,
            j.party_name_snap,
            j.parcel_type,
            j.total_cases,
            f"{j.total_weight:.2f} KG" if j.total_weight else "0.00 KG",
            j.envelope_size,
            j.printer_name,
            j.status
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
